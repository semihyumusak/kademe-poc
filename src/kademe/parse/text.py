"""Ham belgelerden (HTML, PDF, DOCX) paragraf listesi çıkarma ve metin normalleştirme."""
from __future__ import annotations

import re
import subprocess
import unicodedata
from collections import Counter
from pathlib import Path

from bs4 import BeautifulSoup

_WS = re.compile(r"\s+")


def clean(s: str) -> str:
    s = unicodedata.normalize("NFC", s)
    s = s.replace(" ", " ").replace("­", "").replace("​", "")
    s = s.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    return _WS.sub(" ", s).strip()


def tr_lower(s: str) -> str:
    return s.replace("I", "ı").replace("İ", "i").lower()


def fingerprint(s: str) -> str:
    """Sürüm karşılaştırması için biçimden bağımsız parmak izi.

    Boşluk, noktalama ve büyük/küçük harf farklarını, ayrıca şapkalı harf
    (â, î, û) kullanım farklarını yok sayar; PDF/HTML kaynak farkından doğan
    yapay sürüm ayrışmalarını önler.
    """
    s = tr_lower(clean(s))
    s = re.sub(r"\[\d{1,3}\]|\(\d{1,2}\s*\)|[-–—_]{5,}", "", s)
    s = s.translate(str.maketrans("âîû", "aiu"))
    return re.sub(r"[^0-9a-zçğıöşü]", "", s)


# --------------------------------------------------------------------------- HTML

def html_paragraphs(path: Path) -> list[str]:
    raw = path.read_bytes()
    try:
        doc = raw.decode("utf-8")
    except UnicodeDecodeError:  # eski mevzuat.gov.tr sayfaları
        doc = raw.decode("cp1254", errors="replace")
    soup = BeautifulSoup(doc, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    out: list[str] = []
    for el in soup.find_all(["p", "td", "li", "h1", "h2", "h3"]):
        # İç içe blokları tekrar saymamak için yalnızca yaprak blokları al.
        if el.find(["p", "td", "li"]):
            continue
        t = clean(el.get_text(" "))
        if t:
            out.append(t)
    return out


# --------------------------------------------------------------------------- PDF

_PARA_START = re.compile(
    r"^(?:(?:GEÇİCİ|Geçici|EK|Ek)\s+)?(?:MADDE|Madde)\s+\d+|^\(\d+\)\s|^[a-zçğıöşü]\)\s|^\d+\)\s|"
    r"^[A-ZÇĞİÖŞÜ][A-ZÇĞİÖŞÜ ]+BÖLÜM\b|^BÖLÜM\b|^\d+(?:\.\d+)*\.\s+[A-ZÇĞİÖŞÜ]{2,}"
)


def pdf_paragraphs(path: Path) -> list[str]:
    try:
        txt = subprocess.run(
            ["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, check=True
        ).stdout.decode("utf-8", "replace")
    except (FileNotFoundError, subprocess.CalledProcessError) as exc:  # pragma: no cover
        raise RuntimeError(f"pdftotext çalıştırılamadı ({path}): {exc}") from exc
    pages = txt.split("\f")
    lines_per_page = [[clean(l) for l in p.splitlines()] for p in pages]
    # Sayfa üst/alt bilgisi: çok sayfada aynen tekrar eden kısa satırlar.
    cnt = Counter(l for page in lines_per_page for l in set(page) if l and len(l) < 80)
    n_pages = max(1, len(pages))
    noise = {l for l, c in cnt.items() if c >= max(3, n_pages // 3)}
    paras: list[str] = []
    buf = ""
    for page in lines_per_page:
        # Sayfa dipnotları ayırıcı çizgiden sonra gelir; madde metnine karışmaması için atılır.
        for i, line in enumerate(page):
            if re.fullmatch(r"[-–—_]{5,}", line.replace(" ", "")):
                page = page[:i]
                break
        for line in page:
            if not line or line in noise or re.fullmatch(r"\d{1,4}|-\s*\d+\s*-|Sayfa \d+.*", line):
                continue
            new_para = (
                not buf
                or _PARA_START.match(line)
                or (buf.endswith((".", ":", ";")) and line[:1].isupper())
            )
            if new_para:
                if buf:
                    paras.append(buf)
                buf = line
            else:
                buf = buf[:-1] + line if buf.endswith("-") and not buf.endswith(" -") else f"{buf} {line}"
    if buf:
        paras.append(buf)
    return paras


# --------------------------------------------------------------------------- DOCX

def docx_paragraphs(path: Path) -> list[str]:
    import docx  # python-docx

    d = docx.Document(str(path))
    out = [clean(p.text) for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = clean(c.text)
                if t and t not in cells:
                    cells.append(t)
            if cells:
                out.append(" | ".join(cells))
    return [p for p in out if p]


def paragraphs(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return pdf_paragraphs(path)
    if ext == ".docx":
        return docx_paragraphs(path)
    return html_paragraphs(path)
